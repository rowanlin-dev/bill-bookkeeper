#!/usr/bin/env python3
"""
数据迁移脚本：将现有Excel数据导入SQLite
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
import sqlite3
import os
from datetime import datetime

# 路径配置（按脚本位置推导项目根目录，不再硬编码绝对路径）
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, 'data')
EXCEL_FILE = os.path.join(DATA_DIR, '个人财务账本_2026.xlsx')
DB_FILE = os.path.join(DATA_DIR, 'ledger.db')

def init_database():
    """初始化SQLite数据库"""
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    
    # 交易流水
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL,
            type TEXT NOT NULL,
            name TEXT,
            category TEXT,
            amount REAL NOT NULL,
            account TEXT DEFAULT '微信',
            note TEXT,
            source TEXT DEFAULT '用户输入',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    # 负债明细
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS liabilities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            amount REAL NOT NULL,
            creditor TEXT,
            due_date TEXT,
            rate TEXT,
            status TEXT DEFAULT '未还',
            note TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    # 固定支出
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS fixed_expenses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            amount REAL NOT NULL,
            cycle TEXT DEFAULT '月付',
            method TEXT DEFAULT '微信',
            note TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    conn.commit()
    return conn

def migrate_transactions(conn, wb):
    """迁移交易流水"""
    cursor = conn.cursor()
    ws = wb['交易流水']
    
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[1] or row[1] == '日期':
            continue
        
        try:
            date = str(row[1])[:10]
            type_ = row[2] or '支出'
            name = row[3] or ''
            category = row[4] or '其他'
            amount = float(row[5]) if row[5] else 0
            account = row[6] or '微信'
            note = row[7] or ''
            source = row[9] or '用户输入'
            
            if amount <= 0:
                continue
            
            cursor.execute('''
                INSERT INTO transactions (date, type, name, category, amount, account, note, source)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ''', (date, type_, name, category, amount, account, note, source))
            count += 1
        except Exception as e:
            print(f"  跳过行: {e}")
    
    conn.commit()
    print(f"  已迁移 {count} 条交易记录")

def migrate_liabilities(conn, wb):
    """迁移负债明细"""
    cursor = conn.cursor()
    ws = wb['负债明细']
    
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[1] or row[1] == '负债项目':
            continue
        if '已删除' in str(row[1]):
            continue
        
        try:
            name = str(row[1])
            amount = float(row[2]) if row[2] else 0
            creditor = row[3] or ''
            due_date = row[4] or '待定'
            rate = row[5] or '0%'
            status = row[6] or '未还'
            note = row[7] or ''
            
            if amount <= 0:
                continue
            
            cursor.execute('''
                INSERT INTO liabilities (name, amount, creditor, due_date, rate, status, note)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (name, amount, creditor, due_date, rate, status, note))
            count += 1
        except Exception as e:
            print(f"  跳过行: {e}")
    
    conn.commit()
    print(f"  已迁移 {count} 条负债记录")

def migrate_fixed_expenses(conn, wb):
    """迁移固定支出"""
    cursor = conn.cursor()
    ws = wb['固定支出']
    
    count = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[1] or row[1] in ['支出项目', '合计']:
            continue
        
        try:
            name = str(row[1])
            amount = float(row[2]) if row[2] else 0
            cycle = row[3] or '月付'
            method = row[4] or '微信'
            note = row[5] or ''
            
            if amount <= 0:
                continue
            
            cursor.execute('''
                INSERT INTO fixed_expenses (name, amount, cycle, method, note)
                VALUES (?, ?, ?, ?, ?)
            ''', (name, amount, cycle, method, note))
            count += 1
        except Exception as e:
            print(f"  跳过行: {e}")
    
    conn.commit()
    print(f"  已迁移 {count} 条固定支出")

def main():
    print("=" * 60)
    print("数据迁移：Excel → SQLite")
    print("=" * 60)
    
    # 检查Excel文件
    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Excel文件不存在: {EXCEL_FILE}")
        return
    
    print(f"\n📁 Excel文件: {EXCEL_FILE}")
    print(f"📁 数据库: {DB_FILE}")
    
    # 确保数据目录存在
    os.makedirs(DATA_DIR, exist_ok=True)
    
    # 如果数据库已存在，检查是否有数据
    if os.path.exists(DB_FILE):
        conn_check = sqlite3.connect(DB_FILE)
        cursor_check = conn_check.cursor()
        cursor_check.execute('SELECT COUNT(*) FROM transactions')
        count = cursor_check.fetchone()[0]
        conn_check.close()
        
        if count > 0:
            print(f"\n⚠️ 数据库已存在且包含 {count} 条记录")
            print("如果需要重新迁移，请先手动删除数据库文件")
            print(f"  数据库路径: {DB_FILE}")
            return
        else:
            # 数据库存在但为空，继续迁移
            os.remove(DB_FILE)
            print(f"\n🗑️ 删除空数据库: {DB_FILE}")
    
    # 初始化数据库
    conn = init_database()
    print("\n✅ 数据库初始化完成")
    
    # 加载Excel
    print("\n📖 读取Excel...")
    wb = openpyxl.load_workbook(EXCEL_FILE)
    print(f"  Sheets: {wb.sheetnames}")
    
    # 迁移数据
    print("\n🔄 迁移交易流水...")
    migrate_transactions(conn, wb)
    
    print("\n🔄 迁移负债明细...")
    migrate_liabilities(conn, wb)
    
    print("\n🔄 迁移固定支出...")
    migrate_fixed_expenses(conn, wb)
    
    # 验证
    cursor = conn.cursor()
    print("\n📊 迁移结果验证:")
    for table in ['transactions', 'liabilities', 'fixed_expenses']:
        cursor.execute(f'SELECT COUNT(*) FROM {table}')
        count = cursor.fetchone()[0]
        print(f"  {table}: {count} 条记录")
    
    conn.close()
    
    print("\n" + "=" * 60)
    print("✅ 数据迁移完成！")
    print("=" * 60)

if __name__ == '__main__':
    main()
